from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
import os

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from .models import (
    LearnerProfile,
    StudyMaterial,
    Result,
)

from .serializers import (
    AdaptLessonSerializer,
    EvaluateAnswerSerializer,
    LearnerProfileSerializer,
    PracticeQuestionsSerializer,
)

from .llm_service import (
    adapt_lesson_with_ai,
    generate_practice_questions_with_ai,
    evaluate_answer_with_ai,
    voice_chat_with_ai,
    generate_learning_visual_with_ai,
    LLMServiceError,
)


def get_profile(user):
    profile, _ = LearnerProfile.objects.get_or_create(
        user=user
    )

    return profile


def add_topic(items, topic):
    topic = (topic or "General").strip()

    if not topic:
        topic = "General"

    result = [
        item
        for item in items
        if item.lower() != topic.lower()
    ]

    result.insert(0, topic)

    return result[:5]


def remove_topic(items, topic):
    topic = (topic or "").strip().lower()

    return [
        item
        for item in items
        if item.lower() != topic
    ]


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def adapt_lesson(request):
    serializer = AdaptLessonSerializer(
        data=request.data
    )

    serializer.is_valid(
        raise_exception=True
    )

    text = serializer.validated_data["text"]

    subject = serializer.validated_data.get(
        "subject",
        "General",
    )

    preferences = serializer.validated_data.get(
        "preferences",
        {},
    )

    try:
        result = adapt_lesson_with_ai(
            text=text,
            preferences=preferences,
        )

    except LLMServiceError as exc:
        return Response(
            {"detail": str(exc)},
            status=status.HTTP_502_BAD_GATEWAY,
        )

    material = StudyMaterial.objects.create(
        user=request.user,
        title=f"{subject} lesson",
        subject=subject,
        original_text=text,
        adapted_text=result["simplified_text"],
        key_points=result["key_points"],
    )

    return Response(
        {
            "material_id": material.id,
            "subject": subject,
            "simplified_text":
                result["simplified_text"],
            "key_points":
                result["key_points"],
        },
        status=status.HTTP_200_OK,
    )


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def practice_questions(request):
    serializer = PracticeQuestionsSerializer(
        data=request.data
    )

    serializer.is_valid(
        raise_exception=True
    )

    adapted_text = serializer.validated_data[
        "adapted_text"
    ]

    subject = serializer.validated_data.get(
        "subject",
        "General",
    )

    try:
        result = generate_practice_questions_with_ai(
            adapted_text=adapted_text,
        )

    except LLMServiceError as exc:
        return Response(
            {"detail": str(exc)},
            status=status.HTTP_502_BAD_GATEWAY,
        )

    return Response(
        {
            "subject": subject,
            "questions": result["questions"],
        },
        status=status.HTTP_200_OK,
    )


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def evaluate_answer(request):
    serializer = EvaluateAnswerSerializer(
        data=request.data
    )

    serializer.is_valid(
        raise_exception=True
    )

    question = serializer.validated_data[
        "question"
    ]

    student_answer = serializer.validated_data[
        "student_answer"
    ]

    reference_answer = serializer.validated_data[
        "reference_answer"
    ]

    subject = serializer.validated_data.get(
        "subject",
        "General",
    )

    try:
        result = evaluate_answer_with_ai(
            question=question,
            student_answer=student_answer,
            reference_answer=reference_answer,
        )

    except LLMServiceError as exc:
        return Response(
            {"detail": str(exc)},
            status=status.HTTP_502_BAD_GATEWAY,
        )

    profile = get_profile(
        request.user
    )

    Result.objects.create(
        user=request.user,
        question=None,
        student_answer=student_answer,
        correct=result["correct"],
        feedback=result["feedback"],
    )

    profile.questions_answered += 1

    if result["correct"]:
        profile.strong_topics = add_topic(
            profile.strong_topics,
            subject,
        )

        profile.weak_topics = remove_topic(
            profile.weak_topics,
            subject,
        )

    else:
        profile.weak_topics = add_topic(
            profile.weak_topics,
            subject,
        )

    profile.save()

    return Response(
        {
            "question": question,
            "correct": result["correct"],
            "feedback": result["feedback"],
            "questions_answered":
                profile.questions_answered,
        },
        status=status.HTTP_200_OK,
    )


@api_view(["GET", "PATCH"])
@permission_classes([IsAuthenticated])
def learner_profile(request):
    profile = get_profile(
        request.user
    )

    if request.method == "GET":
        serializer = LearnerProfileSerializer(
            profile
        )

        return Response(
            serializer.data,
            status=status.HTTP_200_OK,
        )

    serializer = LearnerProfileSerializer(
        profile,
        data=request.data,
        partial=True,
    )

    serializer.is_valid(
        raise_exception=True
    )

    serializer.save()

    return Response(
        serializer.data,
        status=status.HTTP_200_OK,
    )


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def dashboard_data(request):
    profile = get_profile(
        request.user
    )

    latest_material = (
        StudyMaterial.objects
        .filter(user=request.user)
        .order_by("-created_at")
        .first()
    )

    results = Result.objects.filter(
        user=request.user
    )

    total_answers = results.count()

    correct_answers = results.filter(
        correct=True
    ).count()

    accuracy = (
        round(
            (
                correct_answers
                / total_answers
            ) * 100
        )
        if total_answers
        else 0
    )

    lessons_created = (
        StudyMaterial.objects
        .filter(user=request.user)
        .count()
    )

    if latest_material:
        recent_learning = {
            "material_id":
                latest_material.id,
            "subject":
                latest_material.subject
                or "General",
            "title":
                latest_material.title
                or "Recent lesson",
            "created_at":
                latest_material.created_at,
        }
    else:
        recent_learning = None

    learner_name = (
        request.user.first_name
        or request.user.username
    )

    return Response(
        {
            "learner_name":
                learner_name,
            "pace":
                profile.pace,
            "questions_answered":
                profile.questions_answered,
            "correct_answers":
                correct_answers,
            "accuracy":
                accuracy,
            "lessons_created":
                lessons_created,
            "strong_topics":
                profile.strong_topics,
            "weak_topics":
                profile.weak_topics,
            "recent_learning":
                recent_learning,
        },
        status=status.HTTP_200_OK,
    )


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def progress_data(request):
    profile = get_profile(
        request.user
    )

    results = Result.objects.filter(
        user=request.user
    )

    total_answers = results.count()

    correct_answers = results.filter(
        correct=True
    ).count()

    incorrect_answers = results.filter(
        correct=False
    ).count()

    accuracy = (
        round(
            (
                correct_answers
                / total_answers
            ) * 100
        )
        if total_answers
        else 0
    )

    lessons_created = (
        StudyMaterial.objects
        .filter(user=request.user)
        .count()
    )

    latest_material = (
        StudyMaterial.objects
        .filter(user=request.user)
        .order_by("-created_at")
        .first()
    )

    latest_subject = (
        latest_material.subject
        if latest_material
        else "No topic yet"
    )

    if accuracy >= 80:
        understanding_label = (
            "Strong understanding"
        )

    elif accuracy >= 60:
        understanding_label = (
            "Building confidence"
        )

    elif total_answers > 0:
        understanding_label = (
            "Keep practising"
        )

    else:
        understanding_label = (
            "Ready to begin"
        )

    if profile.weak_topics:
        reflection = (
            f"Keep practising "
            f"{profile.weak_topics[0]}. "
            f"You are building understanding "
            f"step by step."
        )

    elif profile.strong_topics:
        reflection = (
            f"You are doing well with "
            f"{profile.strong_topics[0]}. "
            f"Keep strengthening what you know."
        )

    else:
        reflection = (
            "Complete a practice session and "
            "Ddiba will begin showing "
            "personalized feedback here."
        )

    return Response(
        {
            "questions_answered":
                profile.questions_answered,
            "correct_answers":
                correct_answers,
            "incorrect_answers":
                incorrect_answers,
            "accuracy":
                accuracy,
            "lessons_created":
                lessons_created,
            "strong_topics":
                profile.strong_topics,
            "weak_topics":
                profile.weak_topics,
            "latest_subject":
                latest_subject,
            "understanding_label":
                understanding_label,
            "reflection":
                reflection,
        },
        status=status.HTTP_200_OK,
    )


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def extract_file(request):
    uploaded_file = request.FILES.get("file")

    if not uploaded_file:
        return Response(
            {
                "detail": "No file was uploaded.",
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    # Limit uploads to 10 MB for now.
    max_size = 10 * 1024 * 1024

    if uploaded_file.size > max_size:
        return Response(
            {
                "detail": "File is too large. Maximum size is 10 MB.",
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    filename = uploaded_file.name

    extension = os.path.splitext(
        filename
    )[1].lower()

    supported_extensions = [
        ".txt",
        ".pdf",
        ".docx",
        ".pptx",
    ]

    if extension not in supported_extensions:
        return Response(
            {
                "detail": (
                    "Unsupported file type. "
                    "Please upload TXT, PDF, DOCX, or PPTX."
                ),
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    try:
        extracted_text = ""

        # -----------------------------
        # TXT
        # -----------------------------
        if extension == ".txt":
            raw_content = uploaded_file.read()

            try:
                extracted_text = raw_content.decode(
                    "utf-8"
                )

            except UnicodeDecodeError:
                extracted_text = raw_content.decode(
                    "latin-1"
                )

        # -----------------------------
        # PDF
        # -----------------------------
        elif extension == ".pdf":
            uploaded_file.seek(0)

            reader = PdfReader(
                uploaded_file
            )

            pages = []

            for page in reader.pages:
                page_text = (
                    page.extract_text()
                    or ""
                )

                if page_text.strip():
                    pages.append(
                        page_text.strip()
                    )

            extracted_text = "\n\n".join(
                pages
            )

        # -----------------------------
        # DOCX
        # -----------------------------
        elif extension == ".docx":
            uploaded_file.seek(0)

            document = Document(
                uploaded_file
            )

            paragraphs = []

            for paragraph in document.paragraphs:
                text = paragraph.text.strip()

                if text:
                    paragraphs.append(
                        text
                    )

            # Also extract text from tables.
            for table in document.tables:
                for row in table.rows:
                    values = []

                    for cell in row.cells:
                        cell_text = (
                            cell.text.strip()
                        )

                        if cell_text:
                            values.append(
                                cell_text
                            )

                    if values:
                        paragraphs.append(
                            " | ".join(values)
                        )

            extracted_text = "\n\n".join(
                paragraphs
            )

        # -----------------------------
        # PPTX
        # -----------------------------
        elif extension == ".pptx":
            uploaded_file.seek(0)

            presentation = Presentation(
                uploaded_file
            )

            slides_text = []

            for slide_number, slide in enumerate(
                presentation.slides,
                start=1,
            ):
                slide_parts = []

                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        shape_text = (
                            shape.text.strip()
                        )

                        if shape_text:
                            slide_parts.append(
                                shape_text
                            )

                if slide_parts:
                    slides_text.append(
                        (
                            f"Slide {slide_number}\n"
                            + "\n".join(
                                slide_parts
                            )
                        )
                    )

            extracted_text = "\n\n".join(
                slides_text
            )

        extracted_text = (
            extracted_text.strip()
        )

        if not extracted_text:
            return Response(
                {
                    "detail": (
                        "Ddiba could not find readable text "
                        "inside this file."
                    ),
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        return Response(
            {
                "filename": filename,
                "file_type": extension,
                "extracted_text":
                    extracted_text,
                "character_count":
                    len(extracted_text),
            },
            status=status.HTTP_200_OK,
        )

    except Exception as exc:
        return Response(
            {
                "detail": (
                    f"Could not read this file: {str(exc)}"
                ),
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

@api_view(["POST"])
@permission_classes([IsAuthenticated])
def voice_chat(request):
    message = (
        request.data.get(
            "message",
            ""
        )
        .strip()
    )

    history = request.data.get(
        "history",
        [],
    )

    if not message:
        return Response(
            {
                "detail":
                    "Please say or type something first."
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    if not isinstance(
        history,
        list,
    ):
        history = []

    history = history[-8:]

    profile = get_profile(
        request.user
    )

    preferences = {
        "short_explanations":
            profile.short_explanations,

        "step_by_step":
            profile.step_by_step,

        "examples":
            profile.examples,

        "larger_text":
            profile.larger_text,

        "more_spacing":
            profile.more_spacing,

        "shorter_paragraphs":
            profile.shorter_paragraphs,

        "highlight_words":
            profile.highlight_words,

        "read_aloud":
            profile.read_aloud,

        "pace":
            profile.pace,
    }

    lower_message = (
        message.lower()
    )

    visual_phrases = [
        "visualise",
        "visualize",
        "draw ",
        "draw me",
        "illustrate",
        "make an image",
        "make a picture",
        "create an image",
        "create a picture",
        "generate an image",
        "generate a picture",
        "show me a picture",
        "show me an image",
        "show me a diagram",
        "make a diagram",
        "create a diagram",
        "generate a diagram",
        "visual representation",
    ]

    wants_visual = any(
        phrase in lower_message
        for phrase in visual_phrases
    )

    if wants_visual:
        try:
            image_base64 = (
                generate_learning_visual_with_ai(
                    request_text=message,
                    conversation_history=history,
                )
            )

        except LLMServiceError as exc:
            return Response(
                {
                    "detail":
                        str(exc),
                },
                status=
                    status.HTTP_502_BAD_GATEWAY,
            )

        return Response(
            {
                "message":
                    message,

                "type":
                    "image",

                "response":
                    "Here is a visual to help you understand it.",

                "image":
                    (
                        "data:image/png;base64,"
                        + image_base64
                    ),

                "read_aloud":
                    profile.read_aloud,

                "pace":
                    profile.pace,
            },
            status=status.HTTP_200_OK,
        )

    try:
        answer = voice_chat_with_ai(
            message=message,
            preferences=preferences,
            conversation_history=history,
        )

    except LLMServiceError as exc:
        return Response(
            {
                "detail":
                    str(exc),
            },
            status=
                status.HTTP_502_BAD_GATEWAY,
        )

    return Response(
        {
            "message":
                message,

            "type":
                "text",

            "response":
                answer,

            "image":
                None,

            "read_aloud":
                profile.read_aloud,

            "pace":
                profile.pace,
        },
        status=status.HTTP_200_OK,
    )
@api_view(["GET"])
def test_api(request):
    return Response(
        {
            "message":
                "Ddiba backend is connected successfully"
        }
    )